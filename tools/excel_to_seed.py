"""Convert the Research Dashboard workbook into the app's seed data (assets/seed.js).

Usage: python3 tools/excel_to_seed.py Research_Dashboard.xlsx android/app/src/main/assets/seed.js
"""
import datetime
import json
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    papers, keys = [], {}

    def add_paper(key, **kw):
        if key in keys:
            return keys[key]
        pid = "p%d" % (len(papers) + 1)
        papers.append(dict(id=pid, key=key, title=kw.get("title", ""), authors=kw.get("authors", ""),
                           type=kw.get("type", "Journal article"), override=kw.get("override", ""),
                           notes=kw.get("notes", "")))
        keys[key] = pid
        return pid

    for r in wb["Papers"].iter_rows(min_row=2, values_only=True):
        key = clean(r[0])
        if key:
            add_paper(key, title=clean(r[1]), authors=clean(r[2]), type=clean(r[3]) or "Journal article",
                      override=clean(r[4]), notes=clean(r[12]))

    subs = []
    for r in wb["Submissions"].iter_rows(min_row=2, values_only=True):
        key, journal = clean(r[0]), clean(r[1])
        if not key and not journal:
            continue
        subs.append(dict(id="s%d" % (len(subs) + 1), paper=add_paper(key), journal=journal, rating=clean(r[2]),
                         status=clean(r[3]) or "Planned", submitted=clean(r[4]), decided=clean(r[5]),
                         notes=clean(r[7]), titleUsed=clean(r[8])))

    dash = wb["Dashboard"]
    seed = dict(version=1, settings=dict(chaseDays=int(dash["M3"].value or 90), lastUpdated=clean(dash["M4"].value)),
                papers=papers, subs=subs, portals={})
    with open(dst, "w", encoding="utf-8") as f:
        f.write("window.RD_SEED = " + json.dumps(seed, ensure_ascii=False, indent=1) + ";\n")
    print("papers=%d submissions=%d" % (len(papers), len(subs)))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
